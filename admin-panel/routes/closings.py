"""Closingan CRUD and export routes."""

from datetime import datetime
import io

from flask import Blueprint, current_app, jsonify, request, send_file

from admin_core.auth import login_required, require_internal_auth
from admin_core.db import get_db


closings_bp = Blueprint('closings', __name__)


def closing_filter_query():
    status = request.args.get('status', '')
    date = request.args.get('date', '')
    source = request.args.get('source', '')

    query = 'SELECT * FROM closings WHERE 1=1'
    params = []
    if status:
        query += ' AND status=?'
        params.append(status)
    if date:
        query += ' AND (created_at LIKE ? OR timestamp LIKE ?)'
        params.extend([date + '%', date + '%'])
    if source:
        query += ' AND source=?'
        params.append(source)
    query += ' ORDER BY id DESC'
    return query, params, status, date


@closings_bp.route('/api/closings', methods=['GET'])
@login_required
def get_closings():
    query, params, _, _ = closing_filter_query()
    with get_db() as conn:
        c = conn.cursor()
        c.execute(query, params)
        return jsonify([dict(row) for row in c.fetchall()])


@closings_bp.route('/api/closings', methods=['POST'])
@require_internal_auth
def create_closing():
    data = request.json
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    with get_db() as conn:
        c = conn.cursor()
        c.execute('''INSERT INTO closings
            (timestamp, created_at, user_wa, raw_text, nama, alamat, telepon, kode_pos,
             berat, harga_non_cod, nilai_cod, produk, kelurahan, qty,
             instruksi, courier, gudang, status, source)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
            (data.get('timestamp', now), data.get('created_at', now),
             data.get('user_wa', ''), data.get('raw_text', ''),
             data.get('nama', ''), data.get('alamat', ''), data.get('telepon', ''),
             data.get('kode_pos', ''), data.get('berat', '1'),
             data.get('harga_non_cod', ''), data.get('nilai_cod', ''),
             data.get('produk', ''), data.get('kelurahan', ''), data.get('qty', '1'),
             data.get('instruksi', ''), data.get('courier', ''), data.get('gudang', ''),
             data.get('status', 'draft'), data.get('source', 'Manual')))
        closing_id = c.lastrowid
    current_app.logger.info(f"Closing created: #{closing_id}")
    return jsonify({'success': True, 'id': closing_id})


@closings_bp.route('/api/closings/<int:closing_id>', methods=['PUT'])
@login_required
def update_closing(closing_id):
    data = request.json
    with get_db() as conn:
        c = conn.cursor()
        c.execute('''UPDATE closings SET
            nama=?, alamat=?, telepon=?, kode_pos=?, berat=?,
            harga_non_cod=?, nilai_cod=?, produk=?, kelurahan=?,
            qty=?, instruksi=?, courier=?, gudang=?, status=?
            WHERE id=?''',
            (data.get('nama',''), data.get('alamat',''), data.get('telepon',''),
             data.get('kode_pos',''), data.get('berat','1'),
             data.get('harga_non_cod',''), data.get('nilai_cod',''),
             data.get('produk',''), data.get('kelurahan',''),
             data.get('qty','1'), data.get('instruksi',''),
             data.get('courier',''), data.get('gudang',''),
             data.get('status','draft'), closing_id))
    return jsonify({'success': True})


@closings_bp.route('/api/closings/<int:closing_id>', methods=['DELETE'])
@login_required
def delete_closing(closing_id):
    with get_db() as conn:
        c = conn.cursor()
        c.execute('DELETE FROM closings WHERE id=?', (closing_id,))
    return jsonify({'success': True})


@closings_bp.route('/api/closings/export', methods=['GET'])
@login_required
def export_closings_xls():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        return jsonify({'success': False, 'error': 'Library openpyxl belum terinstall. Jalankan: pip install openpyxl'}), 500

    query, params, status, date = closing_filter_query()
    with get_db() as conn:
        c = conn.cursor()
        c.execute(query, params)
        rows = c.fetchall()

    if not rows:
        return jsonify({'success': False, 'error': 'Tidak ada data untuk di-export'}), 404

    wb = Workbook()
    ws = wb.active
    ws.title = "Closingan"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    headers = ['ID', 'Timestamp', 'Created At', 'Source', 'User WA', 'Nama', 'Telepon',
               'Alamat', 'Kelurahan', 'Kode Pos', 'Produk', 'Qty', 'Berat',
               'Harga Non-COD', 'Nilai COD', 'Courier', 'Gudang',
               'Instruksi', 'Status']
    ws.append(headers)

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for row in rows:
        ws.append([
            row['id'], row['timestamp'], row['created_at'], row['source'], row['user_wa'],
            row['nama'], row['telepon'], row['alamat'], row['kelurahan'], row['kode_pos'],
            row['produk'], row['qty'], row['berat'], row['harga_non_cod'], row['nilai_cod'],
            row['courier'], row['gudang'], row['instruksi'], row['status']
        ])

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 50)

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"sukumba-closings-{datetime.now().strftime('%Y-%m-%d-%H%M%S')}"
    if date:
        filename += f"-{date}"
    if status:
        filename += f"-{status}"
    filename += ".xlsx"

    current_app.logger.info(f"Export XLS: {len(rows)} rows exported")
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )
